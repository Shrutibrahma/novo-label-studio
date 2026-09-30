"""Spreadsheet parsing (spec 10.1): CSV (charset detection, sniffed delimiter), XLSX (openpyxl, formula
results), XLS (xlrd). Spreadsheets are only read, never executed."""

from __future__ import annotations

import csv
import io
import posixpath
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import charset_normalizer
import openpyxl
import xlrd

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_DATA_ROWS = 50_000
KINDS = {".csv": "csv", ".xlsx": "xlsx", ".xls": "xls"}
CSV_DELIMITERS = [",", ";", "\t"]


class ParseError(Exception):
    """`code` is a 14.2 code: FILE_TOO_LARGE, FILE_UNREADABLE, FILE_NO_HEADER, FILE_UNSUPPORTED."""

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


@dataclass
class Sheet:
    name: str
    headers: list[str] = field(default_factory=list)
    rows: list[tuple[int, list[Any]]] = field(default_factory=list)  # (source row number, cells)


def file_kind(filename: str) -> str:
    kind = KINDS.get(Path(filename).suffix.lower())
    if kind is None:
        raise ParseError("FILE_UNSUPPORTED")
    return kind


# ---------------------------------------------------------------- cell normalization
def normalize_cell(value: Any) -> Any:
    """Strings trimmed; dates -> ISO YYYY-MM-DD; numbers kept as numbers; blanks -> None."""
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        return text or None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat(timespec="seconds")
    if isinstance(value, int | float | Decimal):
        return value
    text = str(value).strip()
    return text or None


def number_as_text(value: int | float | Decimal) -> str:
    """Integral numbers without decimals and never in scientific notation (10.1, part_number column)."""
    if isinstance(value, int):
        return str(value)
    d = Decimal(str(value))
    if d == d.to_integral_value():
        return str(int(d))
    return format(d.normalize(), "f")


def _is_empty(cells: list[Any]) -> bool:
    return all(c is None for c in cells)


def _finish(name: str, raw_rows: list[tuple[int, list[Any]]]) -> Sheet:
    """Header row = first non-empty row; fully empty rows are skipped and not counted."""
    non_empty = [(n, cells) for n, cells in raw_rows if not _is_empty(cells)]
    if not non_empty:
        raise ParseError("FILE_NO_HEADER")
    header_row, header_cells = non_empty[0]
    width = max(len(c) for _, c in non_empty)

    def cell(cells: list[Any], col: int) -> Any:
        return cells[col] if col < len(cells) else None

    # Trailing columns with no header and no data are dropped.
    while width > 0 and cell(header_cells, width - 1) is None and all(
            cell(c, width - 1) is None for _, c in non_empty[1:]):
        width -= 1
    if width == 0:
        raise ParseError("FILE_NO_HEADER")
    headers: list[str] = []
    seen: dict[str, int] = {}
    for i in range(width):
        value = cell(header_cells, i)
        h = number_as_text(value) if isinstance(value, int | float | Decimal) else (str(value) if value is not None else "")
        h = h.strip() or f"Column {i + 1}"
        count = seen.get(h.lower(), 0) + 1
        seen[h.lower()] = count
        headers.append(h if count == 1 else f"{h} ({count})")
    data = [(n, (cells + [None] * width)[:width]) for n, cells in non_empty[1:]]
    if len(data) > MAX_DATA_ROWS:
        raise ParseError("FILE_TOO_LARGE")
    return Sheet(name=name, headers=headers, rows=data)


# ---------------------------------------------------------------- CSV
def _decode(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        best = charset_normalizer.from_bytes(data).best()
        if best is None:
            raise ParseError("FILE_UNREADABLE") from None
        return str(best)


def _sniff_delimiter(text: str) -> str:
    sample = "\n".join(text.splitlines()[:50])
    try:
        return csv.Sniffer().sniff(sample, delimiters="".join(CSV_DELIMITERS)).delimiter
    except csv.Error:
        first = next((line for line in text.splitlines() if line.strip()), "")
        counts = {d: first.count(d) for d in CSV_DELIMITERS}
        return max(counts, key=lambda d: counts[d]) if max(counts.values()) > 0 else ","


def read_csv(data: bytes) -> list[Sheet]:
    text = _decode(data)
    if "\x00" in text:
        raise ParseError("FILE_UNREADABLE")
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=_sniff_delimiter(text), quotechar='"')
    rows: list[tuple[int, list[Any]]] = []
    try:
        for n, record in enumerate(reader, start=1):
            cells = [normalize_cell(c) for c in record]
            if _is_empty(cells):
                continue
            rows.append((n, cells))
            if len(rows) > MAX_DATA_ROWS + 1:  # + 1 for the header row
                raise ParseError("FILE_TOO_LARGE")
    except csv.Error as exc:
        raise ParseError("FILE_UNREADABLE") from exc
    return [_finish("", rows)]


# ---------------------------------------------------------------- XLSX pictures
class CellImage:
    """A picture in a cell: Excel's "Place in Cell" picture, or a floating picture whose top-left corner is
    anchored in the cell. Holds the raw image bytes; shown as "[picture]" in samples."""

    __slots__ = ("data",)

    def __init__(self, data: bytes) -> None:
        self.data = data

    def __str__(self) -> str:
        return "[picture]"

    def __repr__(self) -> str:
        return f"CellImage({len(self.data)} bytes)"


_NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "rd": "http://schemas.microsoft.com/office/spreadsheetml/2017/richdata",
    "rvr": "http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel",
    "xlrd": "http://schemas.microsoft.com/office/spreadsheetml/2017/richdata",
}
MAX_PICTURE_BYTES = 10 * 1024 * 1024
_CELL_REF = re.compile(r"^([A-Z]{1,3})(\d+)$")


def _col_index(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def _resolve(base_part: str, target: str) -> str:
    """Relationship target -> zip member name (targets are relative to the part's folder, or absolute)."""
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_part), target))


def _rels(z: zipfile.ZipFile, part: str) -> dict[str, str]:
    rels_name = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    if rels_name not in z.namelist():
        return {}
    root = ElementTree.fromstring(z.read(rels_name))
    return {r.get("Id", ""): _resolve(part, r.get("Target", "")) for r in root.findall("pr:Relationship", _NS)}


def _read_member(z: zipfile.ZipFile, name: str) -> bytes | None:
    try:
        info = z.getinfo(name)
    except KeyError:
        return None
    if info.file_size > MAX_PICTURE_BYTES:
        return None
    return z.read(info)


def _sheet_part(z: zipfile.ZipFile, sheet_name: str) -> str | None:
    wb = ElementTree.fromstring(z.read("xl/workbook.xml"))
    rels = _rels(z, "xl/workbook.xml")
    for s in wb.findall("m:sheets/m:sheet", _NS):
        if s.get("name") == sheet_name:
            return rels.get(s.get(f"{{{_NS['r']}}}id", ""))
    return None


def _rich_value_images(z: zipfile.ZipFile) -> list[bytes | None]:
    """Value-metadata index (the cell's 1-based `vm`, minus one) -> picture bytes, for "Place in Cell" pictures."""
    names = set(z.namelist())
    if "xl/metadata.xml" not in names or "xl/richData/rdrichvalue.xml" not in names:
        return []
    meta = ElementTree.fromstring(z.read("xl/metadata.xml"))
    types = [t.get("name") for t in meta.findall("m:metadataTypes/m:metadataType", _NS)]
    future = {f.get("name"): [b.find(".//xlrd:rvb", _NS) for b in f.findall("m:bk", _NS)]
              for f in meta.findall("m:futureMetadata", _NS)}
    rich_blocks = future.get("XLRICHVALUE", [])
    # Which position in each rich value holds the image relationship index, per structure.
    rel_pos: list[int | None] = []
    if "xl/richData/rdrichvaluestructure.xml" in names:
        for s in ElementTree.fromstring(z.read("xl/richData/rdrichvaluestructure.xml")).findall("rd:s", _NS):
            keys = [k.get("n") for k in s.findall("rd:k", _NS)]
            rel_pos.append(keys.index("_rvRel:LocalImageIdentifier") if "_rvRel:LocalImageIdentifier" in keys else None)
    values = ElementTree.fromstring(z.read("xl/richData/rdrichvalue.xml")).findall("rd:rv", _NS)
    rel_ids: list[str] = []
    if "xl/richData/richValueRel.xml" in names:
        rel_ids = [r.get(f"{{{_NS['r']}}}id", "")
                   for r in ElementTree.fromstring(z.read("xl/richData/richValueRel.xml")).findall("rvr:rel", _NS)]
    targets = _rels(z, "xl/richData/richValueRel.xml")

    def picture(rv_index: int) -> bytes | None:
        if not 0 <= rv_index < len(values):
            return None
        rv = values[rv_index]
        s = int(rv.get("s", "0"))
        pos = rel_pos[s] if s < len(rel_pos) else None
        vs = rv.findall("rd:v", _NS)
        if pos is None or pos >= len(vs) or not (vs[pos].text or "").strip().isdigit():
            return None
        rel = int((vs[pos].text or "").strip())
        if rel >= len(rel_ids) or rel_ids[rel] not in targets:
            return None
        return _read_member(z, targets[rel_ids[rel]])

    out: list[bytes | None] = []
    for bk in meta.findall("m:valueMetadata/m:bk", _NS):
        rc = bk.find("m:rc", _NS)
        data = None
        if rc is not None:
            t, v = int(rc.get("t", "0")), int(rc.get("v", "-1"))
            if 1 <= t <= len(types) and types[t - 1] == "XLRICHVALUE" and 0 <= v < len(rich_blocks):
                rvb = rich_blocks[v]
                if rvb is not None and (rvb.get("i") or "").isdigit():
                    data = picture(int(rvb.get("i") or "0"))
        out.append(data)
    return out


def xlsx_pictures(path: Path, sheet_name: str) -> dict[tuple[int, int], bytes]:
    """(source row number, 0-based column) -> picture bytes for one sheet. Reads only the package XML and media;
    nothing in the workbook is executed. Unreadable picture parts are skipped (the cell then has no picture)."""
    out: dict[tuple[int, int], bytes] = {}
    try:
        with zipfile.ZipFile(path) as z:
            part = _sheet_part(z, sheet_name)
            if part is None or part not in z.namelist():
                return out
            # Floating pictures: anchored at their top-left cell.
            for target in [t for t in _rels(z, part).values() if "/drawings/" in t and t in z.namelist()]:
                drawing = ElementTree.fromstring(z.read(target))
                media = _rels(z, target)
                for anchor in list(drawing.findall("xdr:twoCellAnchor", _NS)) + list(drawing.findall("xdr:oneCellAnchor", _NS)):
                    frm, blip = anchor.find("xdr:from", _NS), anchor.find(".//a:blip", _NS)
                    if frm is None or blip is None:
                        continue
                    col, row = frm.findtext("xdr:col", "", _NS), frm.findtext("xdr:row", "", _NS)
                    embed = blip.get(f"{{{_NS['r']}}}embed", "")
                    if col.isdigit() and row.isdigit() and embed in media:
                        data = _read_member(z, media[embed])
                        if data:
                            out.setdefault((int(row) + 1, int(col)), data)
            # "Place in Cell" pictures: cells with a value-metadata index (vm) that points at a rich value image.
            rich = _rich_value_images(z)
            if rich:
                with z.open(part) as f:
                    for _, el in ElementTree.iterparse(f):
                        if el.tag == f"{{{_NS['m']}}}c":
                            vm, m = el.get("vm"), _CELL_REF.match(el.get("r", ""))
                            if vm and vm.isdigit() and m and 1 <= int(vm) <= len(rich) and rich[int(vm) - 1]:
                                out[(int(m.group(2)), _col_index(m.group(1)))] = rich[int(vm) - 1]  # type: ignore[assignment]
                            el.clear()
    except (zipfile.BadZipFile, ElementTree.ParseError, KeyError, ValueError):
        return {}
    return out


# ---------------------------------------------------------------- XLSX
def read_xlsx(path: Path, only: str | None = None) -> list[Sheet]:
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for corrupt files
        raise ParseError("FILE_UNREADABLE") from exc
    sheets: list[Sheet] = []
    try:
        for ws in wb.worksheets:
            if only is not None and ws.title != only:
                continue
            rows: list[tuple[int, list[Any]]] = []
            non_empty = 0
            pictures = xlsx_pictures(path, ws.title)
            by_row: dict[int, list[tuple[int, bytes]]] = {}
            for (r, c), data in pictures.items():
                by_row.setdefault(r, []).append((c, data))
            for n, record in enumerate(ws.iter_rows(values_only=True), start=1):
                cells = [normalize_cell(c) for c in record]
                for c, data in by_row.get(n, ()):
                    cells += [None] * (c + 1 - len(cells))
                    cells[c] = CellImage(data)  # replaces the "#VALUE!" Excel stores in picture cells
                if _is_empty(cells):
                    continue
                rows.append((n, cells))
                non_empty += 1
                if non_empty > MAX_DATA_ROWS + 1:  # + 1 for the header row
                    raise ParseError("FILE_TOO_LARGE")
            try:
                sheets.append(_finish(ws.title, rows))
            except ParseError as err:
                if err.code != "FILE_NO_HEADER":
                    raise
    finally:
        wb.close()
    return sheets


# ---------------------------------------------------------------- XLS
def _xls_value(book: xlrd.book.Book, cell: xlrd.sheet.Cell) -> Any:
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return normalize_cell(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
        except (xlrd.xldate.XLDateError, ValueError, OverflowError):
            return cell.value
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if cell.value else "FALSE"
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return None
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        v = float(cell.value)
        return int(v) if v.is_integer() else v
    return normalize_cell(cell.value)


def read_xls(data: bytes, only: str | None = None) -> list[Sheet]:
    try:
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
    except Exception as exc:
        raise ParseError("FILE_UNREADABLE") from exc
    sheets: list[Sheet] = []
    try:
        for name in book.sheet_names():
            if only is not None and name != only:
                continue
            sh = book.sheet_by_name(name)
            if sh.nrows > MAX_DATA_ROWS + 10_000:
                raise ParseError("FILE_TOO_LARGE")
            rows = [(r + 1, [_xls_value(book, sh.cell(r, c)) for c in range(sh.ncols)]) for r in range(sh.nrows)]
            try:
                sheets.append(_finish(name, rows))
            except ParseError as err:
                if err.code != "FILE_NO_HEADER":
                    raise
    finally:
        book.release_resources()
    return sheets


def read_file(path: Path, kind: str, only_sheet: str | None = None) -> list[Sheet]:
    """All non-empty sheets (or just `only_sheet`). Raises ParseError."""
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ParseError("FILE_TOO_LARGE")
    if kind == "csv":
        sheets = read_csv(path.read_bytes())
    elif kind == "xlsx":
        sheets = read_xlsx(path, only_sheet)
    elif kind == "xls":
        sheets = read_xls(path.read_bytes(), only_sheet)
    else:
        raise ParseError("FILE_UNSUPPORTED")
    if not sheets:
        raise ParseError("FILE_NO_HEADER")
    return sheets
