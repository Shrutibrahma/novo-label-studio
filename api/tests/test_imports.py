"""M3 gate: import tests (16.1) and acceptance checks A6, A7."""

from __future__ import annotations

import asyncio
import io
import re
import uuid
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

import httpx
import openpyxl
import pytest
import xlwt
from openpyxl.drawing.image import Image as XlImage
from PIL import Image
from sqlalchemy import select, text

from app.db import sessionmaker
from app.imports import service
from app.imports.mapping import normalize_header, signature, suggest
from app.imports.parse import CellImage, read_csv, read_xlsx
from app.models import AuditLog, ImportBatch, Part, PartAlias

V = "/api/v1"


def csv_bytes(rows: list[list[Any]], delimiter: str = ",", encoding: str = "utf-8") -> bytes:
    out = io.StringIO()
    for r in rows:
        out.write(delimiter.join("" if v is None else str(v) for v in r) + "\r\n")
    return out.getvalue().encode(encoding)


def xlsx_bytes(sheets: dict[str, list[list[Any]]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xls_bytes(rows: list[list[Any]]) -> bytes:
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Parts")
    date_style = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            if isinstance(v, date):
                ws.write(r, c, datetime(v.year, v.month, v.day), date_style)
            elif v is not None:
                ws.write(r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def upload(c: httpx.AsyncClient, name: str, data: bytes, wait: bool = True) -> dict[str, Any]:
    r = await c.post(f"{V}/imports", files={"file": (name, data, "application/octet-stream")})
    assert r.status_code == 201, r.text
    return await settle(c, r.json()["id"]) if wait else r.json()


async def settle(c: httpx.AsyncClient, batch_id: str) -> dict[str, Any]:
    for _ in range(600):
        await service.wait_idle(uuid.UUID(batch_id))
        body = (await c.get(f"{V}/imports/{batch_id}")).json()
        if body["status"] not in ("parsing", "validating"):
            return body
        await asyncio.sleep(0.02)
    raise AssertionError("import did not settle")


async def map_and_stage(c: httpx.AsyncClient, batch: dict[str, Any], mapping: dict[str, str | None] | None = None,
                        save_as: str | None = None) -> dict[str, Any]:
    r = await c.put(f"{V}/imports/{batch['id']}/mapping", json={"mapping": mapping or batch["mapping"],
                                                                 "save_as": save_as})
    assert r.status_code == 200, r.text
    staged = await settle(c, batch["id"])
    assert staged["status"] == "staged", staged
    return staged


async def rows(c: httpx.AsyncClient, batch_id: str, action: str | None = None) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    cursor = None
    while True:
        params = {"limit": 200, **({"action": action} if action else {}), **({"cursor": cursor} if cursor else {})}
        page = (await c.get(f"{V}/imports/{batch_id}/rows", params=params)).json()
        out += page["items"]
        cursor = page["next_cursor"]
        if not cursor:
            return out


# ---------------------------------------------------------------- unit: mapping
def test_synonyms_map_correctly() -> None:
    headers = ["PART_NO", "Short Description", "Desc1", "Rev Level", "Supplier", "Item", "Colour"]
    got = suggest(headers, [("colour", "Colour"), ("supplier_code", "Supplier")])
    assert got == {"PART_NO": "part_number", "Short Description": "part_name", "Desc1": "description",
                   "Rev Level": "revision", "Supplier": "supplier_code", "Item": None, "Colour": "colour"}
    for header, target in [("Part #", "part_number"), ("NOVO PN", "part_number"), ("User No.", "part_number"),
                           ("Item Name", "part_name"), ("Title", "part_name"), ("Long Description", "description"),
                           ("REVISION LEVEL", "revision")]:
        assert suggest([header], [])[header] == target, header


def test_header_normalization_and_signature() -> None:
    assert normalize_header("  Part No. ") == "part_no"
    assert normalize_header("Rev -- Level") == "rev_level"
    assert signature(["B", "Part No"]) == signature(["part_no", "b"])
    assert signature(["A"]) != signature(["B"])


def test_csv_encoding_and_delimiter() -> None:
    latin = csv_bytes([["Part No", "Name"], ["NP-1", "Unterlegscheibe groß"]], delimiter=";", encoding="cp1252")
    sheet = read_csv(latin)[0]
    assert sheet.headers == ["Part No", "Name"] and sheet.rows == [(2, ["NP-1", "Unterlegscheibe groß"])]
    tab = csv_bytes([["pn", "name"], [" 00123 ", ' "Quoted, name" '], [None, None], ["X", "Y"]], delimiter="\t")
    sheet = read_csv(b"\xef\xbb\xbf" + tab)[0]
    assert sheet.rows == [(2, ["00123", '"Quoted, name"']), (4, ["X", "Y"])]  # empty row skipped, not counted


# ---------------------------------------------------------------- API flow
async def test_first_import_and_reimport(admin: httpx.AsyncClient) -> None:
    await admin.post(f"{V}/custom-fields", json={"label": "Max torque", "data_type": "number"})
    await admin.post(f"{V}/custom-fields", json={"label": "Released", "data_type": "date"})
    data = xlsx_bytes({"Parts": [
        ["PART_NO", "Name", "Description", "Rev", "Max torque", "Released", "Notes"],
        [10421, "Bearing Housing", "Cast iron", "C", 12.5, date(2026, 3, 1), "ignored"],
        [None, None, None, None, None, None, None],
        ["NP-2", "Shaft", None, "A", "1,200", "3/4/2026", None],
        ["NP-3", None, "No name", "A", None, None, None],
        ["NP-4", "Torque bad", None, "A", "lots", None, None],
        ["np-2", "Dup of NP-2", None, None, None, None, None],
    ]})
    batch = await upload(admin, "NovoParts.xlsx", data)
    assert batch["status"] == "mapping" and batch["save_as_default"] == "NovoParts"
    assert batch["mapping"] == {"PART_NO": "part_number", "Name": "part_name", "Description": "description",
                                "Rev": "revision", "Max torque": "max_torque", "Released": "released", "Notes": None}
    assert [c["samples"] for c in batch["columns"]][0] == ["10421", "NP-2", "NP-3"]
    r = await admin.put(f"{V}/imports/{batch['id']}/mapping", json={"mapping": batch["mapping"] | {"Name": None}})
    assert r.status_code == 422 and r.json()["error"]["message"] == "Map a column to Part Number and Part Name to continue."
    r = await admin.put(f"{V}/imports/{batch['id']}/mapping",
                        json={"mapping": batch["mapping"] | {"Notes": "part_name"}})
    assert r.json()["error"]["message"] == "Each field can be mapped from one column only."

    staged = await map_and_stage(admin, batch, save_as="Novo master")
    assert staged["counts"] == {"total_rows": 5, "new": 1, "updated": 0, "unchanged": 0, "invalid": 4, "missing": 0}
    invalid = {r["source_row"]: [e["msg"] for e in r["errors"]] for r in await rows(admin, batch["id"], "invalid")}
    assert invalid == {
        4: ["Duplicate part number in this file (rows 4, 7)."],
        5: ["Part name is missing."],
        6: ["Max torque must be a number."],
        7: ["Duplicate part number in this file (rows 4, 7)."],
    }
    new = (await rows(admin, batch["id"], "new"))[0]
    assert new["data"] == {"part_number": "10421", "part_name": "Bearing Housing", "description": "Cast iron",
                           "revision": "C", "max_torque": 12.5, "released": "2026-03-01"}

    # Fix one duplicate: both rows become valid (the other stops being a duplicate).
    dup = next(r for r in await rows(admin, batch["id"], "invalid") if r["source_row"] == 7)
    r = await admin.patch(f"{V}/imports/{batch['id']}/rows/{dup['id']}", json={"data": {"part_number": "NP-7"}})
    assert r.json()["action"] == "new"
    counts = (await admin.get(f"{V}/imports/{batch['id']}")).json()["counts"]
    assert counts["new"] == 3 and counts["invalid"] == 2
    two = next(r for r in await rows(admin, batch["id"]) if r["source_row"] == 4)
    assert two["action"] == "new" and two["data"]["max_torque"] == 1200 and two["data"]["released"] == "2026-03-04"

    r = await admin.post(f"{V}/imports/{batch['id']}/commit")
    assert r.status_code == 200, r.text
    assert (r.json()["new"], r.json()["updated"], r.json()["marked_inactive"], r.json()["skipped"]) == (3, 0, 0, 2)
    async with sessionmaker()() as db:
        parts = {p.part_number: p for p in (await db.execute(select(Part))).scalars()}
        assert set(parts) == {"10421", "NP-2", "NP-7"}
        assert parts["10421"].source == "import" and str(parts["10421"].last_import_batch_id) == batch["id"]
        assert (await db.execute(select(AuditLog.action).where(AuditLog.action == "import.commit"))).scalar() == "import.commit"
        admin_id = parts["10421"].created_by
        db.add(PartAlias(part_id=parts["10421"].id, alias="10-32 x 1/2 16", is_label_name=True, created_by=admin_id))
        await db.commit()

    # Re-import: saved mapping by signature; revision change, empty cell never clears, missing part unticked.
    data2 = xlsx_bytes({"Sheet1": [
        ["PART_NO", "Name", "Description", "Rev", "Max torque", "Released", "Notes"],
        ["10421", "Bearing Housing", None, "D", None, None, None],
        ["NP-2", "Shaft", None, "A", 1200, "2026-03-04", None],
    ]})
    again = await upload(admin, "NovoParts-v2.xlsx", data2)
    assert again["saved_mapping_name"] == "Novo master"
    staged = await map_and_stage(admin, again)
    assert staged["counts"]["updated"] == 1 and staged["counts"]["unchanged"] == 1 and staged["counts"]["missing"] == 1
    upd = (await rows(admin, again["id"], "update"))[0]
    assert upd["diff"] == {"revision": ["C", "D"]} and upd["accepted"] is True
    miss = (await rows(admin, again["id"], "missing"))[0]
    assert miss["data"]["part_number"] == "NP-7" and miss["accepted"] is False
    assert staged["accepted_to_import"] == 1
    r = await admin.post(f"{V}/imports/{again['id']}/commit")
    assert (r.json()["updated"], r.json()["marked_inactive"]) == (1, 0)
    async with sessionmaker()() as db:
        p = (await db.execute(select(Part).where(Part.part_number == "10421"))).scalar_one()
        assert p.revision == "D" and p.description == "Cast iron" and p.custom_data["max_torque"] == 12.5
        label = (await db.execute(select(PartAlias.alias).where(PartAlias.part_id == p.id))).scalar_one()
        assert label == "10-32 x 1/2 16"  # imports never touch label names


async def test_a7_revision_change_marks_printed_part_out_of_date(admin: httpx.AsyncClient) -> None:
    """A7: a printed part whose revision changes shows Out of date; its label name stays."""
    await upload_and_commit(admin, [["Part No", "Name", "Rev"], ["NP-10421", "Bearing Housing", "C"],
                                    ["NP-9", "Other", "A"]])
    async with sessionmaker()() as db:
        part = (await db.execute(select(Part).where(Part.part_number == "NP-10421"))).scalar_one()
        other = (await db.execute(select(Part).where(Part.part_number == "NP-9"))).scalar_one()
        db.add(PartAlias(part_id=part.id, alias="10-32 x 1/2 16", is_label_name=True, created_by=part.created_by))
        await db.flush()
        await simulate_printed(db, part, {"part_number": "NP-10421", "revision": "C", "label_name": "10-32 x 1/2 16"})
        await simulate_printed(db, other, {"part_number": "NP-9"})
        await db.commit()
    states = {i["part_number"]: i["label_state"] for i in (await admin.get(f"{V}/parts")).json()["items"]}
    assert states == {"NP-10421": "current", "NP-9": "current"}
    await upload_and_commit(admin, [["Part No", "Name", "Rev"], ["NP-10421", "Bearing Housing", "D"],
                                    ["NP-9", "Other", "B"]])
    items = {i["part_number"]: i for i in (await admin.get(f"{V}/parts")).json()["items"]}
    assert items["NP-10421"]["label_state"] == "out_of_date" and items["NP-10421"]["label_name"] == "10-32 x 1/2 16"
    assert items["NP-9"]["label_state"] == "current"  # revision isn't printed on NP-9's label
    ood = (await admin.get(f"{V}/parts", params={"label_state": "out_of_date"})).json()["items"]
    assert [i["part_number"] for i in ood] == ["NP-10421"]


async def simulate_printed(db: Any, part: Part, printed_fields: dict[str, Any]) -> None:
    """Stands in for a real print (built in M5): one printed_label with the given snapshot.part."""
    import json

    ids = (await db.execute(text(
        "SELECT (SELECT id FROM label_config WHERE scope='default' AND is_current) AS cfg, "
        "(SELECT id FROM printer LIMIT 1) AS printer"))).one()
    req = (await db.execute(text("INSERT INTO print_request (idempotency_key, kind, request_body, created_by) "
                                 "VALUES (gen_random_uuid(), 'print', '{}', :u) RETURNING id"), {"u": part.created_by})).scalar()
    job = (await db.execute(text("INSERT INTO print_job (request_id, seq_in_request, printer_id, kind, payload_zpl, "
                                 "status, created_by) VALUES (:r, 1, :p, 'print', '\\x00', 'queued', :u) RETURNING id"),
                            {"r": req, "p": ids.printer, "u": part.created_by})).scalar()
    await db.execute(text("INSERT INTO printed_label (job_id, seq_in_job, part_id, label_config_id, snapshot, dpi, "
                          "bitmap_png, bitmap_sha256) VALUES (:j, 1, :p, :c, CAST(:s AS jsonb), 203, '\\x00', '\\x00')"),
                     {"j": job, "p": part.id, "c": ids.cfg, "s": json.dumps({"part": printed_fields})})


async def upload_and_commit(c: httpx.AsyncClient, table: list[list[Any]], name: str = "parts.csv") -> dict[str, Any]:
    batch = await upload(c, name, csv_bytes(table))
    await map_and_stage(c, batch)
    r = await c.post(f"{V}/imports/{batch['id']}/commit")
    assert r.status_code == 200, r.text
    return r.json()


async def test_a6_thousand_rows_ten_bad(admin: httpx.AsyncClient) -> None:
    """A6: 1,000-row file with 10 bad rows: 990 imported, 10 listed with exact messages."""
    table: list[list[Any]] = [["Item No", "Item Name", "Revision"]]
    expected: dict[int, str] = {}
    for i in range(1000):
        row = [f"A6-{i:04d}", f"Part {i}", "A"]
        if i % 100 == 7:
            kind = (i // 100) % 3
            if kind == 0:
                row[0] = ""
                expected[i + 2] = "Part number is missing."
            elif kind == 1:
                row[1] = ""
                expected[i + 2] = "Part name is missing."
            else:
                row[2] = "R" * 17
                expected[i + 2] = "Revision is longer than 16 characters."
        table.append(row)
    batch = await upload(admin, "a6.csv", csv_bytes(table))
    staged = await map_and_stage(admin, batch)
    assert staged["counts"]["new"] == 990 and staged["counts"]["invalid"] == 10
    got = {r["source_row"]: r["errors"][0]["msg"] for r in await rows(admin, batch["id"], "invalid")}
    assert got == expected
    result = (await admin.post(f"{V}/imports/{batch['id']}/commit")).json()
    assert (result["new"], result["skipped"]) == (990, 10)
    async with sessionmaker()() as db:
        assert (await db.execute(text("SELECT count(*) FROM part"))).scalar() == 990


async def test_archived_match_and_inactive(admin: httpx.AsyncClient) -> None:
    await upload_and_commit(admin, [["pn", "name"], ["K-1", "Keep"], ["K-2", "Archived later"], ["K-3", "Gone"]])
    k2 = (await admin.get(f"{V}/parts", params={"q": "K-2"})).json()["items"][0]
    await admin.post(f"{V}/parts/{k2['id']}/archive")
    batch = await upload(admin, "p.csv", csv_bytes([["pn", "name"], ["K-1", "Keep"], ["K-2", "Archived later"]]))
    await map_and_stage(admin, batch)
    upd = (await rows(admin, batch["id"], "update"))[0]
    assert upd["diff"] == {"status": ["archived", "active"]} and upd["accepted"] is False
    miss = (await rows(admin, batch["id"], "missing"))[0]
    await admin.patch(f"{V}/imports/{batch['id']}/rows/{miss['id']}", json={"accepted": True})
    await admin.patch(f"{V}/imports/{batch['id']}/rows/{upd['id']}", json={"accepted": True})
    result = (await admin.post(f"{V}/imports/{batch['id']}/commit")).json()
    assert (result["updated"], result["marked_inactive"]) == (1, 1)
    statuses = {i["part_number"]: i["status"] for i in (await admin.get(f"{V}/parts", params={"status": "all"})).json()["items"]}
    assert statuses == {"K-1": "active", "K-2": "active", "K-3": "inactive"}


async def test_commit_is_all_or_nothing(admin: httpx.AsyncClient) -> None:
    batch = await upload(admin, "p.csv", csv_bytes([["pn", "name"], ["T-1", "One"], ["T-2", "Two"], ["T-3", "Three"]]))
    await map_and_stage(admin, batch)
    # Someone creates T-3 manually after staging: the whole commit must be refused and nothing written.
    await admin.post(f"{V}/parts", json={"part_number": "T-3", "part_name": "Manual"})
    r = await admin.post(f"{V}/imports/{batch['id']}/commit")
    assert r.status_code == 409 and r.json()["error"]["code"] == "STALE_WRITE"
    async with sessionmaker()() as db:
        numbers = set((await db.execute(select(Part.part_number))).scalars())
        assert numbers == {"T-3"}
        assert (await db.get(ImportBatch, uuid.UUID(batch["id"]))).status == "staged"  # type: ignore[union-attr]


async def test_new_part_number_equal_to_alias_is_invalid(admin: httpx.AsyncClient) -> None:
    p = (await admin.post(f"{V}/parts", json={"part_number": "HW-1", "part_name": "Washer"})).json()
    await admin.post(f"{V}/parts/{p['id']}/aliases", json={"alias": "WASHER-M6"})
    batch = await upload(admin, "p.csv", csv_bytes([["pn", "name"], ["washer-m6", "Clash"]]))
    await map_and_stage(admin, batch)
    bad = (await rows(admin, batch["id"], "invalid"))[0]
    assert bad["errors"] == [{"field": "part_number",
                              "msg": "This part number is already used as a label name for part HW-1."}]


async def test_multi_sheet_and_xls(admin: httpx.AsyncClient) -> None:
    data = xlsx_bytes({"Readme": [["Notes"], ["hello"]], "Empty": [], "Parts": [["pn", "name"], ["S-1", "One"]]})
    batch = await upload(admin, "multi.xlsx", data)
    assert batch["status"] == "needs_sheet"
    assert batch["sheets"] == [{"name": "Readme", "rows": 1}, {"name": "Parts", "rows": 1}]
    r = await admin.put(f"{V}/imports/{batch['id']}/sheet", json={"sheet_name": "Parts"})
    assert r.status_code == 200
    batch = await settle(admin, batch["id"])
    assert batch["status"] == "mapping" and batch["sheet_name"] == "Parts"
    staged = await map_and_stage(admin, batch)
    assert staged["counts"]["new"] == 1

    xls = xls_bytes([["Part Number", "Part Name", "Released"], [10421.0, "Housing", date(2026, 1, 2)], ["X-9", "Nine", None]])
    await admin.post(f"{V}/custom-fields", json={"label": "Released", "data_type": "date"})
    b2 = await upload(admin, "old.xls", xls)
    assert b2["status"] == "mapping", b2
    staged = await map_and_stage(admin, b2)
    new = {r["data"]["part_number"]: r["data"] for r in await rows(admin, b2["id"], "new")}
    assert new["10421"] == {"part_number": "10421", "part_name": "Housing", "released": "2026-01-02"}


async def test_file_errors(admin: httpx.AsyncClient) -> None:
    r = await admin.post(f"{V}/imports", files={"file": ("parts.pdf", b"%PDF-1.4", "application/pdf")})
    assert r.status_code == 415 and r.json()["error"]["message"] == "Upload a CSV, XLSX or XLS file."
    r = await admin.post(f"{V}/imports", files={"file": ("big.csv", b"a" * (20 * 1024 * 1024 + 1), "text/csv")})
    assert r.status_code == 413 and r.json()["error"]["message"] == "This file is too large. Limit: 20 MB and 50,000 rows."
    b = await upload(admin, "blank.csv", b"\r\n , \r\n")
    assert b["status"] == "failed" and b["error"]["code"] == "FILE_NO_HEADER"
    b = await upload(admin, "broken.xlsx", b"PK\x03\x04 definitely not a workbook")
    assert b["error"]["message"] == "We couldn't read this file. Save it again as CSV or XLSX and retry."
    many = csv_bytes([["pn", "name"]] + [[f"R{i}", "n"] for i in range(50_001)])
    b = await upload(admin, "many.csv", many)
    assert b["error"]["code"] == "FILE_TOO_LARGE"


async def test_expiry_and_discard(admin: httpx.AsyncClient) -> None:
    batch = await upload(admin, "p.csv", csv_bytes([["pn", "name"], ["E-1", "One"]]))
    await map_and_stage(admin, batch)
    async with sessionmaker()() as db:
        await db.execute(text("UPDATE import_batch SET created_at = now() - interval '25 hours'"))
        await db.commit()
    r = await admin.post(f"{V}/imports/{batch['id']}/commit")
    assert r.status_code == 410 and r.json()["error"]["message"] == "This import expired. Upload the file again."
    async with sessionmaker()() as db:
        await service.expire_batches(db)
        await db.commit()
        assert (await db.get(ImportBatch, uuid.UUID(batch["id"]))).status == "discarded"  # type: ignore[union-attr]
    other = await upload(admin, "q.csv", csv_bytes([["pn", "name"], ["E-2", "Two"]]))
    r = await admin.post(f"{V}/imports/{other['id']}/discard")
    assert r.json()["status"] == "discarded"
    assert (await admin.get(f"{V}/imports/{other['id']}")).status_code == 410


async def test_imports_are_admin_only(admin: httpx.AsyncClient, operator: httpx.AsyncClient) -> None:
    batch = await upload(admin, "p.csv", csv_bytes([["pn", "name"], ["O-1", "One"]]))
    checks = [("GET", f"/imports/{batch['id']}", None), ("PUT", f"/imports/{batch['id']}/mapping", {"mapping": {}}),
              ("GET", f"/imports/{batch['id']}/rows", None), ("POST", f"/imports/{batch['id']}/commit", None),
              ("POST", f"/imports/{batch['id']}/discard", None), ("PUT", f"/imports/{batch['id']}/sheet", {"sheet_name": ""}),
              ("PATCH", f"/imports/{batch['id']}/rows/1", {"accepted": False})]
    for method, path, body in checks:
        assert (await operator.request(method, f"{V}{path}", json=body)).status_code == 403, path
    r = await operator.post(f"{V}/imports", files={"file": ("p.csv", b"pn,name\r\nA,B\r\n", "text/csv")})
    assert r.status_code == 403


# ---------------------------------------------------------------- pictures in cells
def png(color: tuple[int, int, int], size: tuple[int, int] = (40, 30)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="PNG")
    return buf.getvalue()


def xlsx_with_cell_pictures(rows: list[list[Any]], pictures: dict[str, bytes]) -> bytes:
    """An XLSX with Excel "Place in Cell" pictures, as Excel 365 writes them: the cell holds #VALUE! with a
    value-metadata index (vm) -> metadata.xml -> rich value -> richValueRel -> media."""
    base = zipfile.ZipFile(io.BytesIO(xlsx_bytes({"Parts": rows})))
    files = {n: base.read(n) for n in base.namelist()}
    sheet = files["xl/worksheets/sheet1.xml"].decode()
    refs = list(pictures)
    for i, ref in enumerate(refs, start=1):
        new = f'<c r="{ref}" t="e" vm="{i}"><v>#VALUE!</v></c>'
        cell = re.search(rf'<c r="{ref}"[^>]*?(/>|>.*?</c>)', sheet)
        if cell:
            sheet = sheet.replace(cell.group(0), new)
        else:  # empty cell: append it to its row
            row = re.sub(r"[A-Z]+", "", ref)
            m = re.search(rf'<row r="{row}"[^>]*>.*?</row>', sheet)
            assert m is not None
            sheet = sheet.replace(m.group(0), m.group(0)[: -len("</row>")] + new + "</row>")
    files["xl/worksheets/sheet1.xml"] = sheet.encode()
    n = len(refs)
    files["xl/metadata.xml"] = (
        '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:xlrd="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata">'
        '<metadataTypes count="1"><metadataType name="XLRICHVALUE"/></metadataTypes>'
        f'<futureMetadata name="XLRICHVALUE" count="{n}">'
        + "".join(f'<bk><extLst><ext uri="x"><xlrd:rvb i="{i}"/></ext></extLst></bk>' for i in range(n))
        + f'</futureMetadata><valueMetadata count="{n}">'
        + "".join(f'<bk><rc t="1" v="{i}"/></bk>' for i in range(n)) + "</valueMetadata></metadata>").encode()
    rd = "http://schemas.microsoft.com/office/spreadsheetml/2017/richdata"
    files["xl/richData/rdrichvaluestructure.xml"] = (
        f'<rvStructures xmlns="{rd}" count="1"><s t="_localImage"><k n="_rvRel:LocalImageIdentifier" t="i"/>'
        '<k n="CalcOrigin" t="i"/></s></rvStructures>').encode()
    files["xl/richData/rdrichvalue.xml"] = (
        f'<rvData xmlns="{rd}" count="{n}">' + "".join(f'<rv s="0"><v>{i}</v><v>5</v></rv>' for i in range(n))
        + "</rvData>").encode()
    files["xl/richData/richValueRel.xml"] = (
        '<richValueRels xmlns="http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        + "".join(f'<rel r:id="rId{i + 1}"/>' for i in range(n)) + "</richValueRels>").encode()
    files["xl/richData/_rels/richValueRel.xml.rels"] = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + "".join(f'<Relationship Id="rId{i + 1}" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                  f'relationships/image" Target="../media/image{i + 1}.png"/>' for i in range(n))
        + "</Relationships>").encode()
    for i, ref in enumerate(refs, start=1):
        files[f"xl/media/image{i}.png"] = pictures[ref]
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return out.getvalue()


def test_reads_place_in_cell_and_floating_pictures(tmp_path: Path) -> None:
    red, blue = png((255, 0, 0)), png((0, 0, 255))
    f = tmp_path / "cell.xlsx"
    f.write_bytes(xlsx_with_cell_pictures([["Part No", "Name", "Photo"], ["P-1", "One", None], ["P-2", "Two", None],
                                           ["P-3", "Three", None]], {"C2": red, "C4": blue}))
    sheet = read_xlsx(f)[0]
    photo = {n: cells[2] for n, cells in sheet.rows}
    assert isinstance(photo[2], CellImage) and photo[2].data == red
    assert photo[3] is None
    assert isinstance(photo[4], CellImage) and photo[4].data == blue

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Part No", "Name", "Photo"])
    ws.append(["F-1", "Floating", None])
    ws.append(["F-2", "No picture", None])
    ws.add_image(XlImage(io.BytesIO(red)), "C2")
    g = tmp_path / "float.xlsx"
    wb.save(g)
    by_row = {n: cells for n, cells in read_xlsx(g)[0].rows}
    assert isinstance(by_row[2][2], CellImage) and by_row[3][2] is None


async def test_import_pictures_become_part_images(admin: httpx.AsyncClient) -> None:
    red, blue, green = png((255, 0, 0)), png((0, 0, 255)), png((0, 160, 0))
    head = ["NOVO P/N", "Part Name", "Photo"]
    data = xlsx_with_cell_pictures([head, ["4705-087", "Lower Plate", None], ["4705-088", "Wheel Spacer", None],
                                    ["4705-089", "No photo", None]], {"C2": red, "C3": red})
    batch = await upload(admin, "photos.xlsx", data)
    assert batch["mapping"]["Photo"] == "image"
    assert [c["samples"] for c in batch["columns"] if c["header"] == "Photo"] == [["[picture]", "[picture]"]]
    staged = await map_and_stage(admin, batch)
    assert staged["counts"]["new"] == 3
    new = {r["data"]["part_number"]: r["data"] for r in await rows(admin, batch["id"], "new")}
    assert new["4705-087"]["image"] == new["4705-088"]["image"]  # identical pictures share one asset
    assert "image" not in new["4705-089"]
    shown = await admin.get(f"{V}/assets/{new['4705-087']['image']}")
    assert shown.status_code == 200 and shown.headers["content-type"] == "image/png"
    assert (await admin.post(f"{V}/imports/{batch['id']}/commit")).status_code == 200
    async with sessionmaker()() as db:
        parts = {p.part_number: p for p in (await db.execute(select(Part))).scalars()}
        assert str(parts["4705-087"].image_asset_id) == new["4705-087"]["image"]
        assert parts["4705-089"].image_asset_id is None and "image" not in parts["4705-087"].custom_data
    first_image = new["4705-087"]["image"]

    # Re-import: a different picture replaces (shown as a change), the same picture is unchanged, an empty cell
    # keeps the existing image, and a new picture fills a part that had none.
    data2 = xlsx_with_cell_pictures([head, ["4705-087", "Lower Plate", None], ["4705-088", "Wheel Spacer", None],
                                     ["4705-089", "No photo", None]], {"C2": blue, "C4": green})
    again = await upload(admin, "photos-v2.xlsx", data2)
    staged = await map_and_stage(admin, again)
    assert (staged["counts"]["updated"], staged["counts"]["unchanged"]) == (2, 1)
    upd = {r["data"]["part_number"]: r["diff"] for r in await rows(admin, again["id"], "update")}
    assert upd["4705-087"]["image"][0] == first_image and upd["4705-087"]["image"][1] != first_image
    assert upd["4705-089"]["image"][0] is None
    # The picture can't be typed over in review.
    row = next(r for r in await rows(admin, again["id"]) if r["data"]["part_number"] == "4705-087")
    r = await admin.patch(f"{V}/imports/{again['id']}/rows/{row['id']}", json={"data": {"image": str(uuid.uuid4())}})
    assert r.json()["data"]["image"] == upd["4705-087"]["image"][1]
    assert (await admin.post(f"{V}/imports/{again['id']}/commit")).status_code == 200
    async with sessionmaker()() as db:
        parts = {p.part_number: p for p in (await db.execute(select(Part))).scalars()}
        assert str(parts["4705-087"].image_asset_id) == upd["4705-087"]["image"][1]
        assert str(parts["4705-088"].image_asset_id) == first_image  # empty cell: kept
        assert parts["4705-089"].image_asset_id is not None


async def test_image_column_without_pictures_is_invalid(admin: httpx.AsyncClient) -> None:
    batch = await upload(admin, "p.csv", csv_bytes([["pn", "name", "photo"], ["T-1", "One", "bolt.jpg"],
                                                    ["T-2", "Two", None]]))
    assert batch["mapping"]["photo"] == "image"
    staged = await map_and_stage(admin, batch)
    assert (staged["counts"]["new"], staged["counts"]["invalid"]) == (1, 1)
    bad = (await rows(admin, batch["id"], "invalid"))[0]
    assert bad["errors"] == [{"field": "image",
                              "msg": "This cell has no picture. Place the picture in the cell in Excel (.xlsx)."}]
    unreadable = xlsx_with_cell_pictures([["pn", "name", "photo"], ["U-1", "One", None]], {"C2": b"not an image"})
    b2 = await map_and_stage(admin, await upload(admin, "u.xlsx", unreadable))
    bad = (await rows(admin, b2["id"], "invalid"))[0]
    assert bad["errors"][0]["msg"] == "The picture in this cell can't be read. Use a PNG, JPEG or WebP picture."


async def test_saved_mapping_from_before_image_import_still_imports_pictures(admin: httpx.AsyncClient) -> None:
    """A mapping saved before image import existed sent the picture column to a text field. The picture column is
    now suggested as the image, and a picture is never stored as text."""
    head = ["NOVO P/N", "Part Name", "Photo"]
    data = xlsx_with_cell_pictures([head, ["S-1", "One", None], ["S-2", "Two", None]], {"C2": png((9, 9, 9))})
    first = await upload(admin, "sample.xlsx", data)
    staged = await map_and_stage(admin, first, first["mapping"] | {"Photo": "description"})  # the old mapping
    assert all("description" not in r["data"] for r in await rows(admin, first["id"], "new"))
    assert (await admin.post(f"{V}/imports/{first['id']}/commit")).status_code == 200

    again = await upload(admin, "sample.xlsx", data)
    assert again["saved_mapping_name"] == "sample" and again["mapping"]["Photo"] == "image"
    staged = await map_and_stage(admin, again)
    upd = (await rows(admin, again["id"], "update"))[0]
    assert upd["data"]["part_number"] == "S-1" and upd["diff"]["image"][0] is None
    assert staged["counts"]["unchanged"] == 1
